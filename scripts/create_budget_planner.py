#!/usr/bin/env python3
"""
Create a sample Financial Budget Planner spreadsheet with complex formulas.
This serves as a generic test dataset for the spreadsheet formula system.
"""

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
from pathlib import Path


def create_budget_planner():
    """Create a comprehensive budget planner with formulas."""

    # Create workbook
    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # Remove default sheet

    # Create sheets
    income_sheet = wb.create_sheet("Income")
    expenses_sheet = wb.create_sheet("Expenses")
    summary_sheet = wb.create_sheet("Summary")

    # Styling
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # ===== INCOME SHEET =====
    income_sheet['A1'] = "Income Category"
    income_sheet['B1'] = "Monthly Amount"
    income_sheet['C1'] = "Annual Amount"

    # Apply header styling
    for cell in ['A1', 'B1', 'C1']:
        income_sheet[cell].fill = header_fill
        income_sheet[cell].font = header_font
        income_sheet[cell].border = border

    # Income categories with sample data
    income_categories = [
        ("Salary", 5000),
        ("Freelance Income", 1200),
        ("Investment Returns", 500),
        ("Rental Income", 800),
        ("Other Income", 200)
    ]

    row = 2
    for category, monthly_amount in income_categories:
        income_sheet[f'A{row}'] = category
        income_sheet[f'B{row}'] = monthly_amount
        # Formula: Annual = Monthly * 12
        income_sheet[f'C{row}'] = f'=B{row}*12'
        income_sheet[f'B{row}'].number_format = '$#,##0.00'
        income_sheet[f'C{row}'].number_format = '$#,##0.00'
        row += 1

    # Totals
    total_row = row
    income_sheet[f'A{total_row}'] = "Total Income"
    income_sheet[f'A{total_row}'].font = Font(bold=True)
    income_sheet[f'B{total_row}'] = f'=SUM(B2:B{total_row-1})'
    income_sheet[f'C{total_row}'] = f'=SUM(C2:C{total_row-1})'
    income_sheet[f'B{total_row}'].number_format = '$#,##0.00'
    income_sheet[f'C{total_row}'].number_format = '$#,##0.00'
    income_sheet[f'B{total_row}'].font = Font(bold=True)
    income_sheet[f'C{total_row}'].font = Font(bold=True)

    # ===== EXPENSES SHEET =====
    expenses_sheet['A1'] = "Expense Category"
    expenses_sheet['B1'] = "Budgeted"
    expenses_sheet['C1'] = "Actual"
    expenses_sheet['D1'] = "Difference"
    expenses_sheet['E1'] = "% of Budget"

    # Apply header styling
    for cell in ['A1', 'B1', 'C1', 'D1', 'E1']:
        expenses_sheet[cell].fill = header_fill
        expenses_sheet[cell].font = header_font
        expenses_sheet[cell].border = border

    # Expense categories
    expense_categories = [
        ("Housing (Rent/Mortgage)", 1500, 1500),
        ("Utilities", 200, 185),
        ("Groceries", 600, 650),
        ("Transportation", 400, 420),
        ("Insurance", 300, 300),
        ("Healthcare", 250, 180),
        ("Entertainment", 300, 350),
        ("Dining Out", 250, 280),
        ("Clothing", 150, 120),
        ("Personal Care", 100, 95),
        ("Education", 200, 200),
        ("Savings", 1000, 1000),
        ("Debt Payments", 500, 500),
        ("Miscellaneous", 200, 175)
    ]

    row = 2
    for category, budgeted, actual in expense_categories:
        expenses_sheet[f'A{row}'] = category
        expenses_sheet[f'B{row}'] = budgeted
        expenses_sheet[f'C{row}'] = actual
        # Formula: Difference = Budgeted - Actual
        expenses_sheet[f'D{row}'] = f'=B{row}-C{row}'
        # Formula: % of Budget = Actual / Budgeted
        expenses_sheet[f'E{row}'] = f'=IF(B{row}=0,0,C{row}/B{row})'

        # Number formatting
        expenses_sheet[f'B{row}'].number_format = '$#,##0.00'
        expenses_sheet[f'C{row}'].number_format = '$#,##0.00'
        expenses_sheet[f'D{row}'].number_format = '$#,##0.00'
        expenses_sheet[f'E{row}'].number_format = '0.0%'
        row += 1

    # Totals
    total_row = row
    expenses_sheet[f'A{total_row}'] = "Total Expenses"
    expenses_sheet[f'A{total_row}'].font = Font(bold=True)
    expenses_sheet[f'B{total_row}'] = f'=SUM(B2:B{total_row-1})'
    expenses_sheet[f'C{total_row}'] = f'=SUM(C2:C{total_row-1})'
    expenses_sheet[f'D{total_row}'] = f'=SUM(D2:D{total_row-1})'
    expenses_sheet[f'E{total_row}'] = f'=IF(B{total_row}=0,0,C{total_row}/B{total_row})'

    for col in ['B', 'C', 'D']:
        expenses_sheet[f'{col}{total_row}'].number_format = '$#,##0.00'
        expenses_sheet[f'{col}{total_row}'].font = Font(bold=True)
    expenses_sheet[f'E{total_row}'].number_format = '0.0%'
    expenses_sheet[f'E{total_row}'].font = Font(bold=True)

    # ===== SUMMARY SHEET =====
    summary_sheet['A1'] = "Financial Summary Dashboard"
    summary_sheet['A1'].font = Font(size=16, bold=True)
    summary_sheet.merge_cells('A1:C1')

    # Summary sections
    row = 3

    # Income Summary
    summary_sheet[f'A{row}'] = "Total Monthly Income"
    summary_sheet[f'B{row}'] = f'=Income!B{len(income_categories)+2}'
    summary_sheet[f'B{row}'].number_format = '$#,##0.00'
    summary_sheet[f'A{row}'].font = Font(bold=True)
    row += 1

    summary_sheet[f'A{row}'] = "Total Annual Income"
    summary_sheet[f'B{row}'] = f'=Income!C{len(income_categories)+2}'
    summary_sheet[f'B{row}'].number_format = '$#,##0.00'
    summary_sheet[f'A{row}'].font = Font(bold=True)
    row += 2

    # Expense Summary
    summary_sheet[f'A{row}'] = "Total Budgeted Expenses"
    summary_sheet[f'B{row}'] = f'=Expenses!B{len(expense_categories)+2}'
    summary_sheet[f'B{row}'].number_format = '$#,##0.00'
    summary_sheet[f'A{row}'].font = Font(bold=True)
    row += 1

    summary_sheet[f'A{row}'] = "Total Actual Expenses"
    summary_sheet[f'B{row}'] = f'=Expenses!C{len(expense_categories)+2}'
    summary_sheet[f'B{row}'].number_format = '$#,##0.00'
    summary_sheet[f'A{row}'].font = Font(bold=True)
    row += 2

    # Net Income
    summary_sheet[f'A{row}'] = "Net Income (Budget)"
    summary_sheet[f'B{row}'] = f'=B3-B6'
    summary_sheet[f'B{row}'].number_format = '$#,##0.00'
    summary_sheet[f'A{row}'].font = Font(bold=True, color="00FF00")
    summary_sheet[f'B{row}'].font = Font(bold=True, color="00FF00")
    row += 1

    summary_sheet[f'A{row}'] = "Net Income (Actual)"
    summary_sheet[f'B{row}'] = f'=B3-B7'
    summary_sheet[f'B{row}'].number_format = '$#,##0.00'
    summary_sheet[f'A{row}'].font = Font(bold=True, color="00FF00")
    summary_sheet[f'B{row}'].font = Font(bold=True, color="00FF00")
    row += 2

    # Savings Rate
    summary_sheet[f'A{row}'] = "Savings Rate"
    summary_sheet[f'B{row}'] = f'=IF(B3=0,0,(B3-B7)/B3)'
    summary_sheet[f'B{row}'].number_format = '0.0%'
    summary_sheet[f'A{row}'].font = Font(bold=True)
    summary_sheet[f'B{row}'].font = Font(bold=True)
    row += 2

    # Budget Performance
    summary_sheet[f'A{row}'] = "Budget Performance"
    summary_sheet[f'A{row}'].font = Font(size=12, bold=True)
    row += 1

    summary_sheet[f'A{row}'] = "Over/Under Budget"
    summary_sheet[f'B{row}'] = f'=Expenses!D{len(expense_categories)+2}'
    summary_sheet[f'B{row}'].number_format = '$#,##0.00'
    row += 1

    summary_sheet[f'A{row}'] = "Budget Adherence"
    summary_sheet[f'B{row}'] = f'=IF(B6=0,0,B7/B6)'
    summary_sheet[f'B{row}'].number_format = '0.0%'

    # Adjust column widths
    for sheet in [income_sheet, expenses_sheet, summary_sheet]:
        sheet.column_dimensions['A'].width = 30
        sheet.column_dimensions['B'].width = 15
        sheet.column_dimensions['C'].width = 15
        if sheet == expenses_sheet:
            sheet.column_dimensions['D'].width = 15
            sheet.column_dimensions['E'].width = 15

    return wb


def main():
    """Create and save the budget planner."""
    print("Creating Financial Budget Planner...")

    wb = create_budget_planner()

    # Save to data directory
    output_dir = Path(__file__).parent.parent / "datasets" / "spreadsheet"
    output_dir.mkdir(parents=True, exist_ok=True)

    output_file = output_dir / "Budget_Planner.xlsx"
    wb.save(output_file)

    print(f"✓ Budget planner created: {output_file}")
    print("\nSpreadsheet contains:")
    print("  - Income sheet with monthly/annual calculations")
    print("  - Expenses sheet with budget vs actual tracking")
    print("  - Summary dashboard with financial metrics")
    print("  - Complex formulas: SUM, IF, percentage calculations")
    print("\nSample natural language queries:")
    print("  - What is my total monthly income?")
    print("  - How much did I spend on groceries?")
    print("  - Am I over or under budget?")
    print("  - What is my savings rate?")
    print("  - Calculate my net income")


if __name__ == "__main__":
    main()
